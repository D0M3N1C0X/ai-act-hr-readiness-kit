"""
Builds deliverables/ai_act_hr_register.xlsx, the client-facing register.

The answers about each AI system are inputs (blue); its class, role, the obligations it triggers and the national
rules on telling the workers are live formulas over those answers, written as the same tests in the same order as
src/register.py. Change an answer and the register reclassifies the system. The Reconciliation sheet holds the
values Python computed and checks each one against the formula that should reproduce it.

Conventions: blue text = input, black = formula, green = link to another sheet, yellow fill = key assumption.
Arial throughout. Only functions that Excel, LibreOffice and Numbers all evaluate.
"""
from datetime import datetime
from pathlib import Path

from excel_twin.determinism import normalise
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

import register
from config import (
    ASSESSMENT_DATE,
    CLASSES,
    COUNTRIES,
    D_ANNEX_III,
    D_CHAPTERS_I_II,
    D_GENERAL,
    DELIVERABLES,
    DEROGATION,
    ENTITIES,
    HIGH_RISK,
    MINIMAL,
    NOT_APPLICABLE,
    OBLIGATIONS,
    PROHIBITED,
    SOURCE_COMMIT,
    SOURCE_REPO,
    TRANSPARENCY,
    WORKER_RULES,
)

OUTPUT = DELIVERABLES / "ai_act_hr_register.xlsx"
REPO = "https://github.com/D0M3N1C0X/ai-act-hr-readiness-kit"

FONT = "Arial"
BLUE, GREEN, INK, MUTED, WHITE = "0000FF", "008000", "1B2430", "5F6B7A", "FFFFFF"
HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
INPUT_HEAD = PatternFill("solid", fgColor="3E5C76")
YELLOW = PatternFill("solid", fgColor="FFFF00")
DATE, COUNT = "d mmm yyyy", "#,##0"
CLASS_FILLS = {PROHIBITED: "F5C6C2", HIGH_RISK: "F8D7B0", DEROGATION: "E4E7EC", TRANSPARENCY: "DCE8F7",
               MINIMAL: "D5ECDC", "Yes": "F8D7B0", "In force": "F5C6C2"}

# The register's input columns, in order, with their headers.
INPUTS = [("system_id", "ID"), ("system", "AI system"), ("function", "HR function"), ("supplier", "Supplier"),
          ("scope", "Department in scope"), ("in_it", "Used in IT"), ("in_pl", "Used in PL"), ("in_de", "Used in DE"),
          ("in_es", "Used in ES"), ("applicants", "Used on applicants"), ("infers_emotions", "Infers emotions"),
          ("medical_or_safety", "For medical or safety reasons"), ("annex_iii", "Annex III point 4"),
          ("profiling", "Profiles people"), ("art6_3_condition", "Art. 6(3) condition"),
          ("interacts", "Talks to people"), ("decides_about_people", "Decides or helps decide about people"),
          ("fully_automated", "Fully automated"), ("monitors_workers", "Monitors workers"),
          ("controls_input_data", "Employer controls the input data"),
          ("on_market_before", "Type and model on the market before 2 Dec 2027")]
CHOICES = {"supplier": '"Vendor,In-house"', "annex_iii": '"None,4(a),4(b)"', "art6_3_condition": '"None,(a),(b),(c),(d)"'}


def font(color=INK, bold=False, italic=False, size=10) -> Font:
    return Font(name=FONT, color=color, bold=bold, italic=italic, size=size)


def put(ws, ref, value, *, color=INK, bold=False, italic=False, size=10, fmt=None, fill=None, wrap=False,
        align=None) -> object:
    cell = ws[ref]
    cell.value = value
    cell.font = font(color, bold, italic, size)
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if wrap or align:
        cell.alignment = Alignment(wrap_text=wrap, horizontal=align, vertical="top")
    return cell


def header(ws, row, labels, start=1, height=30, fill=HEADER_FILL) -> None:
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=label)
        c.font = font(WHITE, bold=True)
        c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = height


def title(ws, text, subtitle=None) -> None:
    put(ws, "A1", text, bold=True, size=14)
    if subtitle:
        put(ws, "A2", subtitle, color=MUTED, italic=True)


def widths(ws, spec) -> None:
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


def colours(ws, rng) -> None:
    for text, colour in CLASS_FILLS.items():
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{text}"'],
                                                      fill=PatternFill("solid", fgColor=colour)))


class Model:
    def __init__(self):
        self.systems = register.load()
        self.python = register.build().set_index("system_id")
        self.summary = register.summary(register.build())
        self.wb = Workbook()
        self.refs = {}
        self.checks = []        # (area, item, python value, workbook ref)

    def check(self, area, item, value, ref) -> None:
        self.checks.append((area, item, value if value != "" else None, ref))

    def build(self, path: Path) -> None:
        names = ["Cover", "Summary", "Register", "Obligations", "Workers", "Timeline", "Headcount", "Settings",
                 "Reconciliation"]
        self.wb.active.title = names[0]
        for n in names[1:]:
            self.wb.create_sheet(n)
        self.settings(self.wb["Settings"])
        self.headcount(self.wb["Headcount"])
        self.register(self.wb["Register"])
        self.obligations(self.wb["Obligations"])
        self.workers(self.wb["Workers"])
        self.timeline(self.wb["Timeline"])
        self.dashboard(self.wb["Summary"])
        self.reconciliation(self.wb["Reconciliation"])
        self.cover(self.wb["Cover"], names)
        for ws in self.wb.worksheets:
            ws.sheet_view.showGridLines = False
            ws.page_setup.orientation = "landscape"
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
        self.wb.calculation.fullCalcOnLoad = True
        stamp = datetime(2026, 1, 1)
        self.wb.properties.creator = "Domenico Perroni"
        self.wb.properties.title = "AI Act Register for HR"
        self.wb.properties.created = stamp
        self.wb.properties.modified = stamp
        path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(path)
        normalise(path)

    # -- inputs ---------------------------------------------------------------------------------------------------------
    def settings(self, ws) -> None:
        title(ws, "Settings", "The dates the register runs on. Blue cells can be changed.")
        widths(ws, {"A": 3, "B": 58, "C": 16, "D": 80})
        rows = [
            ("assessment", "Date of this assessment", ASSESSMENT_DATE, "Days to each deadline count from here."),
            ("d_i_ii", "Articles 4 and 5 apply from", D_CHAPTERS_I_II, "Art. 113(a)."),
            ("d_general", "Article 50 and the rest apply from", D_GENERAL, "Art. 113, first paragraph."),
            ("d_annex", "Articles 6 to 27 apply to Annex III systems from", D_ANNEX_III,
             "Art. 113(c)(i) as amended by Regulation (EU) 2026/1744 (was 2 August 2026)."),
        ]
        for r, (key, label, value, note) in enumerate(rows, start=4):
            put(ws, f"B{r}", label)
            put(ws, f"C{r}", value, color=BLUE, fmt=DATE, fill=YELLOW if key == "d_annex" else None)
            put(ws, f"D{r}", note, color=MUTED, italic=True)
            self.refs[key] = f"Settings!$C${r}"

    def headcount(self, ws) -> None:
        title(ws, "Headcount by employer and department",
              f"Workers active on 30 June 2026, from the payroll extract of {SOURCE_REPO} (commit {SOURCE_COMMIT}).")
        widths(ws, {"A": 10, "B": 20, "C": 12})
        header(ws, 4, ["Employer", "Department", "Workers"], height=18)
        hc = register.headcount()
        for r, row in enumerate(hc.itertuples(index=False), start=5):
            put(ws, f"A{r}", row.country, color=BLUE)
            put(ws, f"B{r}", row.department, color=BLUE)
            put(ws, f"C{r}", int(row.workers), color=BLUE, fmt=COUNT)
        last = 4 + len(hc)
        self.hc = {k: f"Headcount!${c}$5:${c}${last}" for k, c in (("country", "A"), ("dept", "B"), ("n", "C"))}
        put(ws, f"B{last + 1}", "Total", bold=True)
        put(ws, f"C{last + 1}", f"=SUM(C5:C{last})", bold=True, fmt=COUNT)
        self.check("Headcount", "Workers in all four employers", self.summary["workers_total"], f"Headcount!C{last + 1}")

    # -- the register ---------------------------------------------------------------------------------------------------
    def register(self, ws) -> None:
        title(ws, "Register: every AI system the HR function uses, and its class under the AI Act",
              "Blue columns are the answers about each system; the grey-headed columns are formulas. Change an answer "
              "and the system is reclassified.")
        heads = [h for _, h in INPUTS] + ["Workers in scope", "Class", "Role", "High-risk duties apply", "Why"]
        header(ws, 4, heads, height=58)
        for i in range(len(INPUTS)):
            ws.cell(row=4, column=1 + i).fill = INPUT_HEAD
        self.col = {k: get_column_letter(i + 1) for i, (k, _) in enumerate(INPUTS)}
        n = len(INPUTS)
        self.col.update({"workers": get_column_letter(n + 1), "class": get_column_letter(n + 2),
                         "role": get_column_letter(n + 3), "timing": get_column_letter(n + 4),
                         "rationale": get_column_letter(n + 5)})
        widths(ws, {get_column_letter(i + 1): 11 for i in range(n)})
        widths(ws, {"A": 6, "B": 30, "C": 16, self.col["workers"]: 11, self.col["class"]: 22, self.col["role"]: 18,
                    self.col["timing"]: 22, self.col["rationale"]: 100})
        yesno = DataValidation(type="list", formula1='"Yes,No"', allow_blank=False)
        ws.add_data_validation(yesno)
        lists = {k: DataValidation(type="list", formula1=v, allow_blank=False) for k, v in CHOICES.items()}
        for dv in lists.values():
            ws.add_data_validation(dv)
        self.rows = {}
        for r, s in enumerate(self.systems.to_dict("records"), start=5):
            self.rows[s["system_id"]] = r
            for key, _ in INPUTS:
                put(ws, f"{self.col[key]}{r}", s[key], color=BLUE, wrap=key == "system")
                if key in lists:
                    lists[key].add(f"{self.col[key]}{r}")
                elif key.startswith(("in_", "applicants", "infers", "medical", "profiling", "interacts", "decides",
                                     "fully", "monitors", "controls", "on_market")):
                    yesno.add(f"{self.col[key]}{r}")
            C = {k: f"${v}{r}" for k, v in self.col.items()}
            by_country = "+".join(f'({self.hc["country"]}="{c}")*({C["in_" + c.lower()]}="Yes")' for c in COUNTRIES)
            put(ws, f"{self.col['workers']}{r}",
                f'=SUMPRODUCT({self.hc["n"]}*({by_country})*((({C["scope"]}="All")+({self.hc["dept"]}={C["scope"]}))>0))',
                fmt=COUNT)
            put(ws, f"{self.col['class']}{r}", self.class_formula(C), bold=True)
            put(ws, f"{self.col['role']}{r}",
                f'=IF(AND({C["supplier"]}="In-house",{C["class"]}="{HIGH_RISK}"),"Provider and deployer","Deployer")')
            put(ws, f"{self.col['timing']}{r}",
                f'=IF({C["class"]}<>"{HIGH_RISK}","",IF({C["on_market_before"]}="Yes",'
                f'"At its next significant design change","From 2 December 2027"))')
            put(ws, f"{self.col['rationale']}{r}", s["rationale"], color=BLUE)
            p = self.python.loc[s["system_id"]]
            for item, key in (("class", "class"), ("role", "role"), ("high-risk timing", "timing"),
                              ("workers in scope", "workers")):
                v = p[key]
                self.check("Register", f"{s['system_id']} {s['system']}: {item}", v.item() if hasattr(v, "item") else v,
                           f"Register!{self.col[key]}{r}")
        self.first, self.last = 5, 4 + len(self.systems)
        colours(ws, f"{self.col['class']}{self.first}:{self.col['class']}{self.last}")
        ws.freeze_panes = "C5"
        ws.auto_filter.ref = f"A4:{self.col['rationale']}{self.last}"

    @staticmethod
    def class_formula(C) -> str:
        return (f'=IF(AND({C["infers_emotions"]}="Yes",{C["medical_or_safety"]}<>"Yes"),"{PROHIBITED}",'
                f'IF({C["annex_iii"]}<>"None",IF({C["profiling"]}="Yes","{HIGH_RISK}",'
                f'IF({C["art6_3_condition"]}<>"None","{DEROGATION}","{HIGH_RISK}")),'
                f'IF({C["interacts"]}="Yes","{TRANSPARENCY}","{MINIMAL}")))')

    def C(self, r) -> dict:
        return {k: f"Register!${v}${r}" for k, v in self.col.items()}

    def obligation_test(self, cond, C) -> str:
        cls, high = C["class"], f'{C["class"]}="{HIGH_RISK}"'
        return {
            "in use": f'{cls}<>"{PROHIBITED}"',
            "prohibited": f'{cls}="{PROHIBITED}"',
            "interacts": f'AND({C["interacts"]}="Yes",{cls}<>"{PROHIBITED}")',
            "high-risk": high,
            "high-risk and input": f'AND({high},{C["controls_input_data"]}="Yes")',
            "high-risk at work": f'AND({high},{C["applicants"]}<>"Yes")',
            "high-risk and decides": f'AND({high},{C["decides_about_people"]}="Yes")',
            "high-risk in-house": f'AND({high},{C["supplier"]}="In-house")',
            "derogation": f'{cls}="{DEROGATION}"',
        }[cond]

    def worker_test(self, cond, country, C) -> str:
        used = f'{C["in_" + country.lower()]}="Yes"'
        alive = f'{C["class"]}<>"{PROHIBITED}"'
        staff = f'{C["applicants"]}<>"Yes"'
        high = f'{C["class"]}="{HIGH_RISK}"'
        test = {
            "fully automated": f'AND({staff},{C["fully_automated"]}="Yes",OR({high},{C["monitors_workers"]}="Yes"))',
            "monitors": f'{C["monitors_workers"]}="Yes"',
            "in the country": "TRUE",
            "processes performance": f'OR({C["monitors_workers"]}="Yes",AND({staff},{high}))',
            "selects": f'AND({C["annex_iii"]}="4(a)",{C["profiling"]}="Yes",{C["decides_about_people"]}="Yes")',
            "affects work": high,
        }[cond]
        return f"AND({used},{alive},{test})"

    def matrix(self, ws, items, test, area, label_row_note) -> None:
        """Systems down, rules across: Yes where the rule applies. Used by the Obligations and Workers sheets."""
        heads = ["ID", "AI system", "Class"] + [f"{k}\n{name}" for k, name, *_ in items]
        header(ws, 4, heads, height=96)
        widths(ws, {"A": 6, "B": 30, "C": 22, **{get_column_letter(4 + i): 15 for i in range(len(items))}})
        for i, it in enumerate(items):
            put(ws, f"{get_column_letter(4 + i)}5", it[label_row_note], color=MUTED, italic=True, wrap=True)
        ws.row_dimensions[5].height = 40
        r0 = 6
        for k, sid in enumerate(self.systems["system_id"], start=0):
            r, src = r0 + k, self.rows[sid]
            C = self.C(src)
            put(ws, f"A{r}", f"=Register!A{src}", color=GREEN)
            put(ws, f"B{r}", f"=Register!B{src}", color=GREEN)
            put(ws, f"C{r}", f"={C['class']}", color=GREEN)
            for i, it in enumerate(items):
                ref = f"{get_column_letter(4 + i)}{r}"
                put(ws, ref, f'=IF({test(it, C)},"Yes","")', align="center")
                self.check(area, f"{sid}: {it[0]}", self.python.loc[sid, it[0]], f"{ws.title}!{ref}")
        last = r0 + len(self.systems) - 1
        put(ws, f"B{last + 1}", "Systems", bold=True)
        for i in range(len(items)):
            L = get_column_letter(4 + i)
            put(ws, f"{L}{last + 1}", f'=COUNTIF({L}{r0}:{L}{last},"Yes")', bold=True, align="center")
        colours(ws, f"C{r0}:{get_column_letter(3 + len(items))}{last}")
        ws.freeze_panes = "D6"
        return r0, last

    def obligations(self, ws) -> None:
        title(ws, "Obligations: what each system triggers, and from when",
              "Row 5 gives the article. The Timeline sheet puts the obligations in date order.")
        self.obl_rows = self.matrix(ws, OBLIGATIONS, lambda it, C: self.obligation_test(it[3], C), "Obligations", 2)
        last = self.obl_rows[1]
        for i, it in enumerate(OBLIGATIONS):
            self.check("Obligations", f"{it[0]}: systems", self.summary["obligations"][it[0]],
                       f"Obligations!{get_column_letter(4 + i)}{last + 1}")
        r = last + 3
        put(ws, f"B{r}", "Considered and not applicable", bold=True)
        for k, (name, art, why) in enumerate(NOT_APPLICABLE, start=r + 1):
            put(ws, f"B{k}", name)
            put(ws, f"C{k}", art)
            put(ws, f"D{k}", why, color=MUTED, italic=True)

    def workers(self, ws) -> None:
        title(ws, "Telling the workers: the national rules each system triggers",
              "Row 5 gives the source. Under the AI Act itself (Art. 26(7)), see obligation O08 on the Obligations sheet.")
        self.wk_rows = self.matrix(ws, WORKER_RULES, lambda it, C: self.worker_test(it[4], it[1], C), "Workers", 3)
        last = self.wk_rows[1]
        for i, it in enumerate(WORKER_RULES):
            self.check("Workers", f"{it[0]}: systems", self.summary["worker_rules"][it[0]],
                       f"Workers!{get_column_letter(4 + i)}{last + 1}")
        r = last + 3
        header(ws, r, ["Rule", "", "Country", "What it asks", "", "", "When"], height=18)
        for k, (key, country, name, _src, _cond, what, when) in enumerate(WORKER_RULES, start=r + 1):
            put(ws, f"A{k}", key)
            put(ws, f"B{k}", name, wrap=True)
            put(ws, f"C{k}", ENTITIES[country])
            put(ws, f"D{k}", what, wrap=True)
            ws.merge_cells(f"D{k}:F{k}")
            put(ws, f"G{k}", when, wrap=True)
            ws.row_dimensions[k].height = 42

    def timeline(self, ws) -> None:
        title(ws, "Timeline: every obligation in date order",
              "In force means the date has passed at the assessment date in Settings.")
        widths(ws, {"A": 7, "B": 54, "C": 20, "D": 14, "E": 12, "F": 12, "G": 12})
        header(ws, 4, ["Key", "Obligation", "Article", "Applies from", "Systems", "Days from now", "Status"])
        date_ref = {D_CHAPTERS_I_II: self.refs["d_i_ii"], D_GENERAL: self.refs["d_general"], D_ANNEX_III: self.refs["d_annex"]}
        order = sorted(range(len(OBLIGATIONS)), key=lambda i: (OBLIGATIONS[i][4], i))
        last_obl = self.obl_rows[1]
        for r, i in enumerate(order, start=5):
            key, name, art, _cond, when, *_ = OBLIGATIONS[i]
            L = get_column_letter(4 + i)
            put(ws, f"A{r}", key)
            put(ws, f"B{r}", name)
            put(ws, f"C{r}", art)
            put(ws, f"D{r}", f"={date_ref[when]}", color=GREEN, fmt=DATE)
            put(ws, f"E{r}", f"=Obligations!{L}{last_obl + 1}", color=GREEN, align="center")
            put(ws, f"F{r}", f"=D{r}-{self.refs['assessment']}", fmt="#,##0;-#,##0")
            put(ws, f"G{r}", f'=IF(F{r}<=0,"In force","Ahead")')
        colours(ws, f"G5:G{4 + len(OBLIGATIONS)}")
        r = 6 + len(OBLIGATIONS)
        put(ws, f"B{r}", "Days to 2 December 2027, when the high-risk duties start", bold=True)
        put(ws, f"F{r}", f"={self.refs['d_annex']}-{self.refs['assessment']}", bold=True, fmt="#,##0")
        self.days_ref = f"Timeline!F{r}"
        self.check("Timeline", "Days to the high-risk duties", (D_ANNEX_III - ASSESSMENT_DATE).days, self.days_ref)

    def dashboard(self, ws) -> None:
        title(ws, "AI Act readiness for HR: summary",
              f"Dry run for an employer in four countries, {self.summary['workers_total']:,} workers. Not legal advice.")
        widths(ws, {"A": 3, "B": 56, "C": 14, "D": 3, "E": 40, "F": 12})
        put(ws, "B4", "AI systems by class", bold=True, size=11)
        cls_rng = f"Register!${self.col['class']}${self.first}:${self.col['class']}${self.last}"
        for k, c in enumerate(CLASSES, start=5):
            put(ws, f"B{k}", c)
            put(ws, f"C{k}", f'=COUNTIF({cls_rng},B{k})', bold=True, align="center")
            self.check("Summary", f"Systems: {c}", self.summary[f"n_{c}"], f"Summary!C{k}")
        colours(ws, f"B5:B{4 + len(CLASSES)}")
        r = 6 + len(CLASSES)
        put(ws, f"B{r}", "Total", bold=True)
        put(ws, f"C{r}", f"=SUM(C5:C{r - 2})", bold=True, align="center")
        put(ws, f"B{r + 1}", "Systems the employer put into service itself and therefore provides (Art. 3(3))")
        role_rng = f"Register!${self.col['role']}${self.first}:${self.col['role']}${self.last}"
        put(ws, f"C{r + 1}", f'=COUNTIF({role_rng},"Provider and deployer")', bold=True, align="center")
        self.check("Summary", "Systems the employer provides", self.summary["provider"], f"Summary!C{r + 1}")
        put(ws, f"B{r + 2}", "Days to the high-risk duties (2 December 2027)")
        put(ws, f"C{r + 2}", f"={self.days_ref}", color=GREEN, bold=True, align="center")
        put(ws, f"B{r + 3}", "High-risk systems already on the market: duties at their next significant design change")
        tim_rng = f"Register!${self.col['timing']}${self.first}:${self.col['timing']}${self.last}"
        put(ws, f"C{r + 3}", f'=COUNTIF({tim_rng},"At its next significant design change")', bold=True, align="center")
        self.check("Summary", "High-risk systems already on the market", self.summary["legacy"], f"Summary!C{r + 3}")
        put(ws, "E4", "Telling the workers: systems per rule", bold=True, size=11)
        last = self.wk_rows[1]
        for k, it in enumerate(WORKER_RULES, start=5):
            put(ws, f"E{k}", f"{it[0]} · {it[2]}")
            put(ws, f"F{k}", f"=Workers!{get_column_letter(4 + WORKER_RULES.index(it))}{last + 1}", color=GREEN,
                align="center")

    def reconciliation(self, ws) -> None:
        title(ws, "Reconciliation: workbook formulas against the Python register",
              "Column D was written by src/build_workbook.py from the Python results; column E is the live formula.")
        widths(ws, {"A": 14, "B": 62, "C": 3, "D": 26, "E": 26, "F": 12, "G": 8})
        header(ws, 6, ["Area", "Item", "", "Python value", "Workbook value", "Difference", "Match"])
        first, last = 7, 6 + len(self.checks)
        # counted on "Yes", so a check that turns into an error counts as a failure, not as a match
        put(ws, "B3", (f'=IF(COUNTIF(G{first}:G{last},"Yes")=COUNTA(B{first}:B{last}),'
                       f'"All "&COUNTA(B{first}:B{last})&" checks match",'
                       f'(COUNTA(B{first}:B{last})-COUNTIF(G{first}:G{last},"Yes"))&" of "&COUNTA(B{first}:B{last})'
                       f'&" checks do not match")'), bold=True, size=12)
        for r, (area, item, value, ref) in enumerate(self.checks, start=first):
            put(ws, f"A{r}", area)
            put(ws, f"B{r}", item)
            put(ws, f"D{r}", value, color=BLUE)
            put(ws, f"E{r}", f"={ref}", color=GREEN)
            put(ws, f"F{r}", f'=IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),E{r}-D{r},"")')
            put(ws, f"G{r}", (f'=IFERROR(IF(AND(D{r}="",E{r}=""),"Yes",IF(AND(ISNUMBER(D{r}),ISNUMBER(E{r})),'
                              f'IF(ABS(E{r}-D{r})<=1E-9*MAX(1,ABS(D{r})),"Yes","No"),IF(EXACT(D{r},E{r}),"Yes","No"))),"No")'))
        for text, colour in (("Yes", "D5ECDC"), ("No", "F5C6C2")):
            ws.conditional_formatting.add(f"G{first}:G{last}", CellIsRule(operator="equal", formula=[f'"{text}"'],
                                                                         fill=PatternFill("solid", fgColor=colour)))
        ws.freeze_panes = "A7"
        self.recon_result = "Reconciliation!B3"

    def cover(self, ws, names) -> None:
        widths(ws, {"A": 3, "B": 24, "C": 100})
        put(ws, "B2", "AI Act Register for HR", bold=True, size=18)
        put(ws, "B3", "Regulation (EU) 2024/1689 as amended by Regulation (EU) 2026/1744 - a dry run for an employer "
                      "in four Member States", color=MUTED, size=12)
        put(ws, "B5", "What it does", bold=True, size=11)
        for i, line in enumerate([
            "Lists every AI system the HR function uses and classifies it: prohibited, high-risk, not high-risk under "
            "Article 6(3), transparency only, or minimal risk.",
            "Says which obligations each system triggers and from when, including where building a system in-house "
            "makes the employer its provider.",
            "Says which national rules on informing and consulting workers apply in Italy, Germany, Spain and Poland.",
        ], start=6):
            put(ws, f"C{i}", line)
        put(ws, "B10", "How to read it", bold=True, size=11)
        for i, (label, meaning, colour, fill) in enumerate([
            ("Blue text", "an input: an answer about a system, a date", BLUE, None),
            ("Black text", "a formula", INK, None),
            ("Green text", "a link to another sheet", GREEN, None),
            ("Yellow fill", "a key assumption that moves the results", INK, YELLOW)], start=11):
            put(ws, f"B{i}", label, color=colour, fill=fill)
            put(ws, f"C{i}", meaning)
        put(ws, "B16", "Sheets", bold=True, size=11)
        purpose = {
            "Summary": "Systems by class, the provider question, days to the high-risk duties, national rules.",
            "Register": "One row per AI system: the answers that decide its class, the class and why.",
            "Obligations": "Which of thirteen obligations each system triggers, with the article.",
            "Workers": "Which national rules on telling the workers each system triggers.",
            "Timeline": "The obligations in date order: what is already in force, what starts in December 2027.",
            "Headcount": "Workers by employer and department, to size each system's scope.",
            "Settings": "The application dates of the AI Act and the date of this assessment.",
            "Reconciliation": "Each workbook figure checked against the Python register.",
        }
        for i, n in enumerate(names[1:], start=17):
            cell = put(ws, f"B{i}", n, color="1F5FA8")
            cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{n}'!A1")
            put(ws, f"C{i}", purpose[n])
        put(ws, "B26", "Check", bold=True, size=11)
        put(ws, "C26", f"={self.recon_result}", color=GREEN, bold=True)
        put(ws, "B28", "Limits", bold=True, size=11)
        put(ws, "C28", "Synthetic organisation and systems. Not legal advice. The classification follows the AI Act "
                       "text; Commission guidelines and national practice can move boundary cases. Sources and their "
                       "status are in docs/verification.md in the repository.", wrap=True)
        ws.row_dimensions[28].height = 40
        put(ws, "B30", "Domenico Perroni", color=MUTED)
        cell = put(ws, "C30", REPO, color="1F5FA8")
        cell.hyperlink = REPO


def build(path: Path = OUTPUT) -> Model:
    m = Model()
    m.build(path)
    return m


def main() -> None:
    m = build()
    print(f"workbook: {len(m.checks)} reconciliation checks -> {OUTPUT.relative_to(OUTPUT.parents[1])}")


if __name__ == "__main__":
    main()
