"""The workbook's formulas reproduce the Python register, and they are live: change an answer, the class follows."""
import pytest

pytest.importorskip("formulas")

from excel_twin import values_from_formulas  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import build_workbook  # noqa: E402
import check_workbook  # noqa: E402


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    path = tmp_path_factory.mktemp("wb") / "register.xlsx"
    return build_workbook.build(path), path


def test_every_reconciliation_row_matches(built):
    m, path = built
    assert check_workbook.check(values_from_formulas(path), len(m.checks)) == 0


def test_changing_an_answer_reclassifies_the_system(built, tmp_path):
    m, path = built
    wb = load_workbook(path)
    ws = wb["Register"]
    r = m.rows["S04"]                                   # the interview scheduler, not high-risk under Art. 6(3)
    ws[f"{m.col['profiling']}{r}"] = "Yes"              # it starts ranking candidates
    changed = tmp_path / "changed.xlsx"
    wb.save(changed)
    values = values_from_formulas(changed)
    assert values[("REGISTER", f"{m.col['class']}{r}")] == "High-risk"
    assert check_workbook.check(values, len(m.checks)) == 1   # and the reconciliation notices
